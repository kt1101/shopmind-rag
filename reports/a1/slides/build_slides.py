"""Build the Vietnamese A1 deck from a shared content model.

Run with Python 3.11+ and python-pptx. HTML embeds fonts already installed on
macOS, and remains usable with its fallback fonts when those files are absent.
Source page numbers refer to physical PDF pages, including the cover/contents.
"""
from pathlib import Path
from html import escape
import base64
import json

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

OUT = Path(__file__).resolve().parent
SOURCE = "Nhom4_BaoCao-A1.pdf"
TEXT_RESULTS = [
    ["BM25", .4625, .3568, .3249, 25.86, 44.82],
    ["Dense text", .0125, .0042, .0018, 51.28, 96.68],
    ["Text → image", .0125, .0036, .0013, 40.56, 59.53],
    ["Hybrid no RRF", .4625, .3568, .3249, 67.04, 102.53],
    ["Hybrid RRF", .3250, .1846, .1414, 65.44, 93.54],
]
SLIDES = [
 dict(kind="cover", section="BÁO CÁO BÀI TẬP 1", title="ShopMind", subtitle="Tìm kiếm sản phẩm đa phương thức", pages="1",
      notes="Kính chào cô và các bạn. Nhóm 4 trình bày ShopMind A1 trong môn Tìm kiếm thông minh, Trường Đại học Công nghệ Thông tin, ĐHQG TP.HCM. GVHD: TS. Phạm Đỗ Kim Chi. Thành viên: Trần Văn Quang – 250104022; Nguyễn Minh Thái – 250201090; Trần Khoa – 250104010; Đinh Ngọc Phúc – 26848170. Báo cáo tháng 9/2026. Bài trình bày tập trung vào tầng truy hồi sản phẩm, kết quả thực nghiệm và giới hạn của benchmark."),
 dict(kind="problem", section="01 / BÀI TOÁN", title="Từ truy vấn đến top-k sản phẩm", pages="3", lead="Đánh giá riêng chất lượng truy hồi trước khi mở rộng hệ thống.",
      cards=[("Đầu vào", "Văn bản hoặc hình ảnh\nVí dụ: “giày đen dùng hằng ngày”"), ("Truy hồi", "BM25 · SigLIP2 · kNN\nKết hợp thứ hạng bằng RRF"), ("Đầu ra", "Danh sách top-k sản phẩm\nID, tiêu đề, ảnh và thứ hạng")],
      takeaway="Phạm vi A1: retrieval; chưa dùng LLM để sinh câu trả lời.",
      notes="Người dùng có thể tìm bằng mô tả văn bản hoặc ảnh. Hệ thống trả về một danh sách sản phẩm đã xếp hạng. A1 tập trung hoàn toàn vào retrieval, để đo được đóng góp của từng thành phần trước khi thêm RAG. Câu truy vấn trên slide lấy từ phần giới hạn của báo cáo, chỉ dùng để minh họa loại truy vấn; slide không gắn cho nó một kết quả thực nghiệm mới."),
 dict(kind="architecture", section="01 / KIẾN TRÚC", title="Hai nhánh truy hồi, một catalog", pages="3–4", lead="Elasticsearch lưu metadata và vector; SigLIP2 tạo biểu diễn văn bản/ảnh.",
      notes="Phần offline bắt đầu từ ABO, chuẩn hóa metadata và ảnh, sau đó tạo embedding SigLIP2 cho văn bản và ảnh. Các bản ghi và vector được đưa vào Elasticsearch. Ở phần online, truy vấn văn bản có thể đi qua BM25, kNN trên text vector hoặc kNN trên image vector; truy vấn ảnh đi qua image encoder và image vector. Hybrid trong A1 kết hợp BM25 với dense text bằng RRF. Các chế độ độc lập cũng trả kết quả qua API. Sơ đồ này tổng hợp kiến trúc mô tả trong báo cáo, không đưa phần A2/A3 vào A1."),
 dict(kind="dataset", section="02 / DỮ LIỆU", title="Tập con ABO có thể tái tạo", pages="3–4", lead="Dataset: abo_a1_1500_v1 · seed 42 · tối đa 150 sản phẩm/category.",
      stats=[("1.500", "sản phẩm"), ("215", "category"), ("1", "ảnh chính / sản phẩm"), ("768", "chiều · chuẩn hóa L2")],
      steps=["Loại bản ghi thiếu ID/title/ảnh", "Xáo trộn và lấy mẫu theo category", "Chỉ giữ ảnh tải thành công", "Chuẩn hóa metadata + search_text"],
      takeaway="Mọi phương pháp dùng chung catalog; ảnh small-image có cạnh lớn tối đa 256 px.",
      notes="ABO gốc có 147.702 product listings và 398.212 catalog images theo công trình được trích trong báo cáo. A1 sử dụng một tập con 1.500 sản phẩm, 215 category để kiểm soát pipeline và tái lập thí nghiệm. Script loại dữ liệu thiếu các trường bắt buộc, lấy mẫu với seed 42, giới hạn theo category và tải dư ứng viên để thay ảnh lỗi. search_text ghép title, description, brand, category và attributes. Quy mô nhỏ hỗ trợ so sánh nội bộ, nhưng chưa cho phép suy rộng latency sang catalog thương mại lớn."),
 dict(kind="benchmark", section="02 / BENCHMARK", title="100 truy vấn, 200 cặp qrels", pages="3–4", lead="Sinh xác định từ 100 sản phẩm đa dạng theo category.",
      cards=[("80 truy vấn text", "45 tiếng Anh · 35 tiếng Việt\nExact · semantic/use-case · hard-negative"), ("20 truy vấn ảnh", "Ảnh catalog → crop → resize\nGiảm contrast → nén lại JPEG")],
      takeaway="Mỗi query: 1 target mức 2 + 1 distractor mức 0. Grade 1 chưa dùng; qrels chưa review độc lập.",
      notes="Benchmark có 80 truy vấn text và 20 truy vấn ảnh. Các ảnh truy vấn đã được biến đổi nên không trùng byte với ảnh trong index, nhưng vẫn cùng vật thể và gần như cùng góc nhìn. Qrels có đúng 200 cặp: một positive grade 2 và một distractor grade 0 cho mỗi query. Đây là seed labels, chưa hoàn tất review độc lập bởi hai thành viên và adjudication. Các metric phù hợp để so sánh cấu hình trong cùng benchmark, chưa phải ground truth đã được con người xác nhận."),
 dict(kind="modes", section="03 / PHƯƠNG PHÁP", title="Bốn đường truy hồi cơ bản", pages="4–5", lead="Encoder: google/siglip2-base-patch16-224 · vector văn bản/ảnh chuẩn hóa L2.",
      cards=[("BM25", "Text → trường metadata\nKhớp từ khóa; boost title, brand/category và ID"), ("Dense text", "Text → SigLIP2 → text_vector\nkNN trong không gian văn bản"), ("Cross-modal", "Text → SigLIP2 → image_vector\nÁnh xạ mô tả sang hình ảnh"), ("Image dense", "Ảnh → SigLIP2 → image_vector\nTìm sản phẩm qua ảnh truy vấn")],
      takeaway="Hybrid dùng hai thành phần BM25 và dense text.",
      notes="BM25 truy hồi lexical trên title, brand, category, description, attributes và search_text. Title được boost cao nhất; phrase match và Product ID exact match cũng được ưu tiên. Dense text mã hóa query rồi tìm trên text_vector. Cross-modal dùng cùng query text nhưng tìm trên image_vector. Image dense mã hóa ảnh người dùng và tìm trên image_vector. SigLIP2 là vision-language encoder tổng quát; hiệu quả zero-shot phụ thuộc vào representation, domain và phân bố query."),
 dict(kind="fusion", section="03 / HYBRID", title="RRF kết hợp thứ hạng", pages="4–5", lead="Tối đa 100 ứng viên từ BM25 và 100 ứng viên từ dense text.",
      formula="score(d) = Σᵢ 1 / (60 + rankᵢ(d))",
      cards=[("Hybrid no RRF · E3", "Giữ thứ tự BM25 trước\nNối các ứng viên dense chưa xuất hiện"), ("Hybrid RRF · E4", "Cộng điểm theo rank của từng nhánh\nXếp hạng lại danh sách ứng viên")],
      takeaway="Dùng rank thay cho việc chuẩn hóa trực tiếp BM25 score và cosine score.",
      notes="RRF cộng nghịch đảo của hằng số k cộng với rank, với k bằng 60 trong A1. Tài liệu vắng ở một danh sách không nhận đóng góp từ danh sách đó. Nhánh hybrid no RRF là cấu hình đối chứng trong đánh giá: giữ BM25 trước, rồi nối các dense candidate chưa có. Cách đối chứng này giúp phân biệt việc mở rộng tập ứng viên với việc thay đổi thứ hạng. Fusion cần tín hiệu thành phần có chất lượng hoặc bổ sung được lỗi cho nhau; kết quả A1 sẽ kiểm tra điều kiện đó."),
 dict(kind="experiments", section="04 / THIẾT KẾ THÍ NGHIỆM", title="Đối chứng từng thành phần", pages="5", lead="candidate_k = 100 · top_k = 10 · tắt reranker để cô lập retrieval/fusion.",
      rows=[("E0", "BM25", "Baseline lexical"), ("E1", "Dense text", "Đóng góp của dense"), ("E2", "Text → image", "Cross-modal alignment"), ("E3", "Hybrid no RRF", "Thêm candidate; giữ rank BM25"), ("E4", "Hybrid RRF", "Tác động của rank fusion"), ("E5", "Image dense", "Truy hồi bằng ảnh · tập riêng")],
      takeaway="Recall@10: tìm thấy target · nDCG@10: chất lượng thứ hạng · MRR@10: vị trí relevant đầu tiên.",
      notes="E0 đến E4 được chạy trên cùng 80 text query; E5 chạy riêng trên 20 image query. Giữ candidate_k và top_k, đồng thời tắt reranker để tập trung vào retrieval và fusion. Recall@10 đo tỷ lệ sản phẩm relevant được đưa vào top 10. Vì mỗi query hiện chỉ có một positive, Recall trung bình cũng là tỷ lệ query có target trong top 10. MRR là trung bình nghịch đảo rank của kết quả relevant đầu tiên, bằng 0 nếu không có trong top 10. nDCG dùng gain theo relevance và chiết khấu theo vị trí, chuẩn hóa theo xếp hạng lý tưởng. Latency dùng p50/p95."),
 dict(kind="results", section="05 / KẾT QUẢ TEXT", title="BM25 đứng đầu trên benchmark hiện tại", pages="5–6", lead="So sánh chất lượng trên cùng 80 truy vấn văn bản.",
      takeaway="BM25 và Hybrid no RRF cùng Recall@10 = 0,4625; RRF giảm còn 0,3250.",
      notes="Bảng giữ nguyên số liệu từ Bảng 3 của báo cáo, làm tròn bốn chữ số thập phân. BM25 đạt Recall 0,4625, nDCG 0,3568 và MRR 0,3249. Dense text và text-to-image đều chỉ đạt Recall 0,0125. Hybrid no RRF giữ nguyên chất lượng BM25; hybrid RRF giảm ở cả ba metric. Các con số chỉ áp dụng cho benchmark hiện tại với qrels dạng seed. Không trộn dòng image dense vào bảng text vì nó dùng một tập truy vấn khác."),
 dict(kind="latency", section="05 / ĐỘ TRỄ", title="Chất lượng tốt hơn, độ trễ thấp hơn", pages="5–6", lead="Text query · latency p50/p95 (ms) · cùng các cấu hình E0–E4.",
      takeaway="Hybrid RRF: p50 65,44 ms — khoảng 2,53 lần BM25 (25,86 ms).",
      notes="Mỗi hàng có hai thanh dùng chung trục bắt đầu từ 0: p50 màu teal và p95 màu xanh nhạt. BM25 có p50 25,86 ms và p95 44,82 ms, thấp nhất trong năm cấu hình. Hybrid no RRF mất 67,04/102,53 ms và không cải thiện chất lượng; hybrid RRF mất 65,44/93,54 ms và giảm chất lượng. Tỷ số 65,44 chia 25,86 xấp xỉ 2,53. Đây là latency ở quy mô tập con A1, không phải cam kết hiệu năng production."),
 dict(kind="language", section="05 / PHÂN TÍCH NGÔN NGỮ", title="Khoảng cách Anh–Việt còn lớn", pages="5–6", lead="BM25 · Recall@10 theo ngôn ngữ truy vấn.",
      bars=[("Tiếng Anh · 45 query", .6444), ("Tiếng Việt · 35 query", .2286)],
      cards=[("Quan sát", "Metadata chủ yếu bằng tiếng Anh;\nquery tiếng Việt từ bảng ánh xạ nhỏ."), ("Diễn giải", "Lexical hiện tương thích tiếng Anh hơn.\nChưa cô lập được nguyên nhân duy nhất.")],
      notes="BM25 đạt Recall@10 0,6444 trên 45 query tiếng Anh và 0,2286 trên 35 query tiếng Việt. Đây là phân tích cùng benchmark chứ không phải hai nghiên cứu độc lập. Báo cáo nhắc rằng canonical metadata chủ yếu là tiếng Anh và query tiếng Việt dùng bảng ánh xạ category/màu tương đối nhỏ. Vì thế không nên quy toàn bộ khoảng cách cho thuật toán BM25. Kết luận thận trọng là pipeline hiện có độ tương thích lexical tốt hơn với tiếng Anh."),
 dict(kind="dense", section="06 / CHẨN ĐOÁN DENSE", title="Dense chỉ tìm thấy target ở 1/80 query", pages="6", lead="Dense text và text → image đều Recall@10 = 0,0125; BM25 tương ứng 37/80.",
      cards=[("Encoder / domain", "SigLIP2 chưa fine-tune riêng\ncho product retrieval trên ABO."), ("Representation", "search_text dài so với query ngắn\ncó thể làm loãng tín hiệu."), ("Tính nhất quán", "Chưa pin model_revision; kiểm tra\npadding và encode đơn lẻ/batch."), ("Qrels", "Một positive/query; semantic neighbor\nhợp lý vẫn có thể không được chấm.")],
      takeaway="Bốn giả thuyết cần kiểm tra; chưa xác nhận một nguyên nhân duy nhất.",
      notes="Recall 0,0125 nhân với 80 bằng 1 query có target trong top 10; BM25 là 37 query. Hai phép quy đổi trên slide được tính từ số liệu và cấu trúc qrels trong báo cáo, không phải kết quả chạy mới. So với BM25, Recall dense giảm xấp xỉ 97,3%. Báo cáo nêu bốn khả năng: chưa thích nghi domain, representation văn bản dài, tính nhất quán model/preprocessing và benchmark chỉ có một positive. Không kết luận lỗi embedding chỉ dựa vào metric; cần ablation và kiểm tra representation trước."),
 dict(kind="ablation", section="06 / PHÂN TÍCH ABLATION", title="Suy giảm xuất hiện khi thay đổi rank", pages="6–7", lead="Recall@10 trên cùng 80 text query.",
      stats=[("0,4625", "E0 · BM25"), ("0,4625", "E3 · Hybrid no RRF"), ("0,3250", "E4 · Hybrid RRF")],
      cards=[("E0 → E3", "Nối candidate dense phía sau:\nkhông đổi ba metric chất lượng."), ("E3 → E4", "Dense tham gia thay đổi thứ hạng:\nRecall giảm tương đối 29,7%.")],
      takeaway="So với BM25, RRF giảm nDCG 48,3% và MRR 56,5% trên benchmark này.",
      notes="E0 và E3 có cùng cả ba metric. Ở phần lớn query, BM25 đã lấp đầy top 10 nên thêm ứng viên phía sau ít có cơ hội ảnh hưởng. Khi RRF cho nhánh dense tác động lên rank, Recall giảm từ 0,4625 xuống 0,3250: giảm 0,1375 điểm tuyệt đối hoặc khoảng 29,7% tương đối. nDCG và MRR giảm khoảng 48,3% và 56,5%. Đây là bằng chứng trong cấu hình A1 rằng một nhánh dense yếu có thể làm nhiễu fusion; không phải kết luận tổng quát rằng RRF luôn kém."),
 dict(kind="image", section="05 / KẾT QUẢ IMAGE", title="1,0 trên ảnh catalog biến đổi nhẹ", pages="5–6", lead="Tập riêng: 20 image query · p50 140,82 ms · p95 230,01 ms.",
      stats=[("1,0000", "Recall@10"), ("1,0000", "nDCG@10"), ("1,0000", "MRR@10")],
      steps=["Ảnh catalog", "Crop + resize", "Contrast + JPEG", "Image kNN"],
      takeaway="Đo độ ổn định với biến đổi nhẹ; cần ảnh chụp độc lập để đánh giá tình huống thực tế.",
      notes="Image dense đạt cả ba metric bằng 1 trên 20 ảnh truy vấn. MRR bằng 1 ở đây cũng cho thấy relevant đầu tiên ở rank 1 trong từng query theo qrels hiện có. Kết quả hỗ trợ nhận định pipeline embedding ảnh ổn định trước crop nhẹ, resize, thay contrast và nén JPEG. Query vẫn đến từ cùng ảnh catalog, gần như cùng góc nhìn và background. Vì vậy đây là thử nghiệm robustness trên transformed catalog images, chưa chứng minh kết quả tương đương trên ảnh điện thoại, nền phức tạp hay góc chụp khác. Không so sánh trực tiếp với bảng text như cùng một tập đo."),
 dict(kind="limits", section="07 / LỖI VÀ GIỚI HẠN", title="Metric cần được đọc cùng giới hạn", pages="7", lead="Các kết quả hiện tại phục vụ chẩn đoán và so sánh nội bộ.",
      cards=[("Qrels chưa review", "Hai reviewer + adjudication còn thiếu.\nChỉ 1 positive/query; grade 1 chưa dùng."), ("20 failure cases", "Nhãn lexical_mismatch / semantic_confusion\nlà gợi ý; reviewer_note còn trống."), ("Độ đại diện", "1.500 sản phẩm; ảnh query phụ thuộc catalog.\nChưa đo ảnh thực tế và tải production.")],
      takeaway="Review query, metadata, ảnh target và top retrieved trước khi kết luận nguyên nhân lỗi.",
      notes="Qrels là giới hạn lớn nhất. Với một truy vấn chung, nhiều sản phẩm có thể hợp lý nhưng không được gắn relevance, đặc biệt bất lợi cho semantic retrieval. Cần review độc lập, adjudication và bổ sung grade 1 khi phù hợp. Ablation runner xuất 20 query có nDCG thấp nhất, nhưng nhãn tự gợi ý và trường reviewer_note trống chỉ tạo danh sách để kiểm tra thủ công. Catalog nhỏ và ảnh truy vấn không độc lập cũng hạn chế độ đại diện. Không mô tả những việc còn thiếu trong báo cáo như đã được hoàn tất."),
 dict(kind="id", section="08 / KIỂM TRA HỒI QUY", title="Tra cứu Product ID đạt 100/100 top-1", pages="7", lead="Kiểm tra riêng ngày 19/09/2026 · 100 Product ID đầu tiên · query chữ thường.",
      stats=[("0/100", "Trước tối ưu"), ("100/100", "Sau tối ưu")],
      cards=[("Thay đổi", "Boost Product ID và khớp\nkhông phân biệt hoa/thường."), ("Ý nghĩa", "Cải thiện exact lookup; metric BM25\ntrên 80 text query không đổi.")],
      takeaway="Kết quả hồi quy định danh được báo cáo riêng với benchmark semantic.",
      notes="Theo mục 8 của PDF, kiểm tra ngày 19/09/2026 lấy 100 Product ID đầu tiên với query chữ thường. Sau Product ID boost và matching không phân biệt hoa/thường, exact ID top-1 tăng từ 0/100 lên 100/100. Đây là số liệu được trích từ báo cáo đã cung cấp, không phải lần chạy lại trong quá trình làm slide. Phải tách nó khỏi benchmark 80 text query vì cải thiện tìm định danh không tương đương cải thiện semantic retrieval; báo cáo cho biết các metric BM25 của benchmark không thay đổi."),
 dict(kind="conclusion", section="09 / KẾT LUẬN VÀ HƯỚNG TIẾP THEO", title="Giữ baseline, củng cố benchmark", pages="7–8", lead="A1 đã có pipeline đa phương thức và khung ablation có kiểm soát.",
      cards=[("01 · Benchmark", "Review qrels, thêm relevance phù hợp\nvà ảnh truy vấn chụp độc lập."), ("02 · Dense", "Pin revision/preprocessing; so sánh\ntitle, metadata ngắn và search_text."), ("03 · Fusion", "Sau khi dense ổn định: thử candidate_k,\nweighted RRF, reranker và phân tích cứu/tụt rank.")],
      takeaway="Hiện tại: BM25 tốt nhất cho text; image tốt trên tập dễ; cải thiện dense trước khi tối ưu fusion.",
      notes="Thông điệp kết luận có ba phần. Thứ nhất, A1 đã thiết lập pipeline retrieval end-to-end và khung đánh giá tách thành phần. Thứ hai, BM25 là text baseline tốt nhất ở benchmark hiện tại; dense yếu và RRF giảm chất lượng, trong khi image cần kiểm chứng bằng ảnh độc lập. Thứ ba, ưu tiên chất lượng benchmark và representation trước. Các ablation representation được báo cáo đề xuất gồm title-only, title + brand + category + attributes, và toàn bộ search_text. Sau khi pin model revision và kiểm tra encode đơn lẻ/batch, mới tiếp tục tối ưu candidate window, fusion và reranker. Những bước này là kế hoạch, chưa phải kết quả đã đo."),
 dict(kind="references", section="TÀI LIỆU THAM KHẢO / Q&A", title="Nguồn và cơ sở phương pháp", pages="9", lead="Slide tổng hợp từ báo cáo Nhóm 4; đánh số tài liệu theo PDF gốc.",
      cards=[("[1], [9] · Dữ liệu ABO", "Collins et al. · CVPR 2022\nAmazon Web Services · Registry of Open Data"), ("[2], [3], [8] · Truy hồi", "Robertson & Zaragoza · BM25, 2009\nKarpukhin et al. · DPR, 2020; Thakur et al. · BEIR, 2021"), ("[4], [5] · Vision–language", "Zhai et al. · SigLIP, ICCV 2023\nTschannen et al. · SigLIP 2, arXiv 2025"), ("[6], [7] · Fusion / đánh giá", "Cormack et al. · RRF, SIGIR 2009\nJärvelin & Kekäläinen · nDCG, TOIS 2002")],
      takeaway="Cảm ơn cô và các bạn. Trao đổi về benchmark, representation và kết quả ablation.",
      notes="Nguồn chính: Nhom4_BaoCao-A1.pdf, Nhóm 4, tháng 9/2026. Các tài liệu sau được giữ theo danh mục của báo cáo, không phải một lượt nghiên cứu web mới: [1] Collins et al., ABO, CVPR 2022, DOI 10.1109/CVPR52688.2022.02045. [2] Robertson & Zaragoza, The Probabilistic Relevance Framework: BM25 and Beyond, 2009, DOI 10.1561/1500000019. [3] Karpukhin et al., Dense Passage Retrieval for Open-Domain Question Answering, EMNLP 2020, DOI 10.18653/v1/2020.emnlp-main.550. [4] Zhai et al., Sigmoid Loss for Language Image Pre-Training, ICCV 2023. [5] Tschannen et al., SigLIP 2, arXiv:2502.14786, 2025. [6] Cormack et al., Reciprocal Rank Fusion, SIGIR 2009, DOI 10.1145/1571941.1572114. [7] Järvelin & Kekäläinen, Cumulated Gain-Based Evaluation of IR Techniques, TOIS 2002, DOI 10.1145/582415.582418. [8] Thakur et al., BEIR, NeurIPS Datasets and Benchmarks 2021. [9] AWS Registry of Open Data, Amazon Berkeley Objects Dataset, ngày truy cập theo báo cáo: 29/09/2026. Xem phần ghi chú nguồn trong README; không dùng slide này để suy rộng trạng thái license."),
]

def esc(s):
    return escape(str(s)).replace("\n", "<br>")

def comma(v, n=4):
    return f"{v:.{n}f}".replace(".", ",")

def cards_html(items):
    return '<div class="cards count-%d">%s</div>' % (len(items), ''.join(
        f'<article class="card"><span class="card-index">{i+1:02}</span><h3>{esc(a)}</h3><p>{esc(b)}</p></article>' for i,(a,b) in enumerate(items)))

def stats_html(items):
    return '<div class="stats">'+''.join(f'<div class="stat"><strong>{esc(a)}</strong><span>{esc(b)}</span></div>' for a,b in items)+'</div>'

def steps_html(items):
    return '<ol class="steps">'+''.join(f'<li><span>{i+1:02}</span>{esc(s)}</li>' for i,s in enumerate(items))+'</ol>'

def body_html(s):
    k=s['kind']
    if k=='architecture':
        return '''<div class="architecture"><div class="pipeline"><span class="lane">OFFLINE</span><div>ABO<br><small>metadata + ảnh</small></div><b>→</b><div>Chuẩn hóa<br><small>catalog chung</small></div><b>→</b><div>SigLIP2<br><small>text + image vectors</small></div><b>→</b><div class="focus">Elasticsearch<br><small>metadata + vectors</small></div></div><div class="online"><span class="lane">ONLINE</span><div class="query">Text / ảnh truy vấn</div><div class="retrievers"><div>BM25 <small>metadata</small></div><div>kNN <small>text_vector / image_vector</small></div></div><div class="query">Độc lập / Hybrid RRF<br><small>Hybrid = BM25 + dense text</small></div><div class="focus">API → top-k sản phẩm</div></div></div>'''
    if k=='results':
        rows=''.join('<tr'+(' class="best"' if i==0 else '')+'><th scope="row">'+esc(r[0])+'</th>'+''.join('<td>'+comma(v)+'</td>' for v in r[1:4])+'</tr>' for i,r in enumerate(TEXT_RESULTS))
        return '<div class="table-wrap"><table><thead><tr><th>Cấu hình</th><th>Recall@10</th><th>nDCG@10</th><th>MRR@10</th></tr></thead><tbody>'+rows+'</tbody></table></div>'
    if k=='experiments':
        return '<div class="experiments">'+''.join(f'<div class="experiment"><b>{a}</b><strong>{esc(b)}</strong><span>{esc(c)}</span></div>' for a,b,c in s['rows'])+'</div>'
    if k=='latency':
        rows=''.join(f'<div class="latency-row"><strong>{esc(r[0])}</strong><div class="bar-pair"><div class="bar p50" style="--value:{r[4]/120*100}%"><span>{comma(r[4],2)}</span></div><div class="bar p95" style="--value:{r[5]/120*100}%"><span>{comma(r[5],2)}</span></div></div></div>' for r in TEXT_RESULTS)
        return '<div class="legend"><span class="p50-key">p50</span><span class="p95-key">p95</span><span>Trục chung 0–120 ms</span></div><div class="latency-chart">'+rows+'</div>'
    if k=='language':
        chart='<div class="language-chart">'+''.join(f'<div class="language-row"><span>{esc(a)}</span><div class="language-track"><div style="width:{v*100}%"></div></div><strong>{comma(v*100,2)}%</strong></div>' for a,v in s['bars'])+'</div>'
        return chart+cards_html(s['cards'])
    result=''
    if 'stats' in s: result+=stats_html(s['stats'])
    if 'formula' in s: result+='<div class="formula">'+esc(s['formula'])+'</div>'
    if 'cards' in s: result+=cards_html(s['cards'])
    if 'steps' in s: result+=steps_html(s['steps'])
    return result

CSS = r'''
/* Viewport-fitting base (frontend-slides skill), themed below. */
html,body{height:100%;overflow-x:hidden}html{scroll-snap-type:y mandatory;scroll-behavior:smooth}
.slide{width:100vw;height:100vh;height:100dvh;overflow:hidden;scroll-snap-align:start;display:flex;flex-direction:column;position:relative}
.slide-content{flex:1;display:flex;flex-direction:column;justify-content:center;max-height:100%;overflow:hidden;padding:var(--slide-padding)}
:root{--title-size:clamp(1.5rem,5vw,4rem);--h2-size:clamp(1.25rem,3.5vw,2.5rem);--h3-size:clamp(1rem,2.5vw,1.75rem);--body-size:clamp(.75rem,1.5vw,1.125rem);--small-size:clamp(.65rem,1vw,.875rem);--slide-padding:clamp(1rem,4vw,4rem);--content-gap:clamp(.5rem,2vw,2rem);--element-gap:clamp(.25rem,1vw,1rem)}
.card,.container,.content-box{max-width:min(90vw,1000px);max-height:min(80vh,700px)}
.feature-list,.bullet-list{gap:clamp(.4rem,1vh,1rem)}.feature-list li,.bullet-list li{font-size:var(--body-size);line-height:1.4}
.grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(min(100%,250px),1fr));gap:clamp(.5rem,1.5vw,1rem)}
img,.image-container{max-width:100%;max-height:min(50vh,400px);object-fit:contain}
@media(max-height:700px){:root{--slide-padding:clamp(.75rem,3vw,2rem);--content-gap:clamp(.4rem,1.5vw,1rem);--title-size:clamp(1.25rem,4.5vw,2.5rem);--h2-size:clamp(1rem,3vw,1.75rem)}}
@media(max-height:600px){:root{--slide-padding:clamp(.5rem,2.5vw,1.5rem);--content-gap:clamp(.3rem,1vw,.75rem);--title-size:clamp(1.1rem,4vw,2rem);--body-size:clamp(.7rem,1.2vw,.95rem)}.nav-dots,.keyboard-hint,.decorative{display:none}}
@media(max-height:500px){:root{--slide-padding:clamp(.4rem,2vw,1rem);--title-size:clamp(1rem,3.5vw,1.5rem);--h2-size:clamp(.9rem,2.5vw,1.25rem);--body-size:clamp(.65rem,1vw,.85rem)}}
@media(max-width:600px){:root{--title-size:clamp(1.25rem,7vw,2.5rem)}.grid{grid-template-columns:1fr}}
@media(prefers-reduced-motion:reduce){*,*::before,*::after{animation-duration:.01ms!important;transition-duration:.2s!important}html{scroll-behavior:auto}}
/* Theme: paper, navy and teal; editorial type with data-first geometry. */
:root{--paper:#F5F3ED;--ink:#152C3C;--muted:#536774;--teal:#086F69;--coral:#B44B34;--line:#D7DFDA;--soft:#E6EDE8;--slide-padding:clamp(1rem,3.8vw,4rem);--body-size:clamp(.875rem,1.5vw,1.375rem);--h2-size:clamp(1.5rem,3.4vw,3rem);--h3-size:clamp(1rem,1.7vw,1.5rem);--small-size:clamp(.7rem,1vw,.95rem)}
*{box-sizing:border-box}body{margin:0;background:var(--paper);color:var(--ink);font-family:DeckBody,Verdana,sans-serif}
h1,h2,h3,p{margin:0}h1,h2{font-family:DeckTitle,Georgia,serif;font-weight:700}h2{font-size:var(--h2-size);line-height:1.18;letter-spacing:-.025em;max-width:28em}
.slide-content{width:min(100%,1600px);margin:auto;justify-content:flex-start;gap:clamp(.75rem,2vh,1.5rem);padding-bottom:clamp(3rem,7vh,5rem)}
.eyebrow{font-size:var(--small-size);letter-spacing:.12em;color:var(--teal);font-weight:bold;padding-bottom:clamp(.3rem,1vh,.8rem);border-bottom:1px solid var(--line)}
.lead{font-size:var(--body-size);color:var(--muted);line-height:1.5;max-width:65em}
.body{flex:1;min-height:0;display:flex;flex-direction:column;justify-content:center;gap:clamp(.75rem,2vh,1.5rem)}
.cards{display:grid;gap:clamp(.6rem,1.4vw,1.5rem);grid-template-columns:repeat(2,minmax(0,1fr))}.count-3{grid-template-columns:repeat(3,minmax(0,1fr))}
.card{background:#FFFDF8;border:1px solid var(--line);padding:clamp(.8rem,2vw,1.8rem);max-height:none;max-width:none;min-width:0;position:relative;border-top:3px solid var(--teal)}
.card-index{font-size:var(--small-size);color:var(--teal);display:block;margin-bottom:clamp(.3rem,1vh,.8rem)}h3{font-size:var(--h3-size);line-height:1.3;margin-bottom:clamp(.5rem,1.2vh,1rem)}.card p{font-size:var(--body-size);line-height:1.6;color:var(--muted);overflow-wrap:anywhere}
.takeaway{font-size:var(--body-size);line-height:1.5;color:var(--teal);padding:clamp(.6rem,1.3vw,1rem) clamp(.7rem,1.6vw,1.4rem);background:var(--soft);border-left:4px solid var(--teal)}
.source{position:absolute;bottom:clamp(.75rem,2vh,1.8rem);left:var(--slide-padding);font-size:var(--small-size);color:var(--muted);max-width:70%}.number{position:absolute;bottom:clamp(.75rem,2vh,1.8rem);right:var(--slide-padding);font-size:var(--small-size);font-variant-numeric:tabular-nums}
.stats{display:grid;grid-auto-flow:column;grid-auto-columns:1fr;gap:clamp(.5rem,2vw,2rem)}.stat{border-top:1px solid var(--line);padding-top:clamp(.7rem,2vh,1.5rem)}.stat strong{display:block;font-family:DeckTitle,Georgia,serif;font-size:clamp(2rem,5vw,5rem);color:var(--teal);line-height:1.1}.stat span{display:block;font-size:var(--body-size);line-height:1.4;color:var(--muted);margin-top:clamp(.5rem,1vh,1rem)}
.steps{padding:0;list-style:none;display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:clamp(.5rem,1vw,1rem);margin:clamp(.5rem,1vh,1rem) 0}.steps li{font-size:var(--body-size);line-height:1.45;background:var(--soft);padding:clamp(.6rem,1.3vw,1rem)}.steps span{display:block;color:var(--teal);font-size:var(--small-size);font-weight:bold;margin-bottom:clamp(.3rem,1vh,.6rem)}
.formula{font-family:DeckTitle,Georgia,serif;font-size:clamp(1.4rem,3vw,3rem);background:var(--ink);color:var(--paper);padding:clamp(1rem,3vh,2rem);text-align:center}
.table-wrap table{width:100%;border-collapse:collapse;font-size:var(--body-size);font-variant-numeric:tabular-nums}.table-wrap th,.table-wrap td{padding:clamp(.65rem,1.8vh,1.25rem) clamp(.5rem,1.2vw,1.1rem);border-bottom:1px solid var(--line);text-align:right}.table-wrap th:first-child{text-align:left}.table-wrap thead{background:var(--ink);color:var(--paper)}.table-wrap tbody th{font-weight:normal}.table-wrap .best{background:var(--soft);color:var(--teal);font-weight:bold}.table-wrap .best th{font-weight:bold}
.experiments{display:grid;gap:clamp(.2rem,.5vh,.5rem)}.experiment{display:grid;grid-template-columns:clamp(2rem,4vw,4rem) 1fr 1.8fr;gap:clamp(.5rem,1vw,1rem);padding:clamp(.5rem,1vh,.8rem) clamp(.7rem,1.2vw,1rem);font-size:var(--body-size);border-bottom:1px solid var(--line);align-items:center}.experiment b{color:var(--teal)}.experiment span{color:var(--muted)}
.legend{display:flex;gap:clamp(.8rem,2vw,2rem);font-size:var(--small-size);color:var(--muted)}.p50-key::before,.p95-key::before{content:'';display:inline-block;width:.7em;height:.7em;margin-right:.5em;background:var(--teal)}.p95-key::before{background:#A9C8C5}.latency-chart{display:grid;gap:clamp(.6rem,1.5vh,1rem)}.latency-row{display:grid;grid-template-columns:clamp(7.4rem,18vw,16rem) 1fr;align-items:center;gap:clamp(.5rem,1vw,1rem);font-size:var(--body-size)}.latency-row strong{font-weight:normal}.bar-pair{display:flex;flex-direction:column;gap:clamp(.16rem,.5vh,.35rem);padding-right:clamp(2.8rem,5vw,5rem);border-left:1px solid var(--line)}.bar{width:var(--value);height:clamp(.65rem,1.6vh,1.2rem);position:relative;background:var(--teal)}.bar.p95{background:#A9C8C5}.bar span{position:absolute;left:100%;top:50%;transform:translateY(-50%);padding-left:clamp(.3rem,.7vw,.7rem);font-size:var(--small-size);font-variant-numeric:tabular-nums;color:var(--ink)}
.language-chart{display:grid;gap:clamp(.8rem,2vh,1.5rem)}.language-row{display:grid;grid-template-columns:1.1fr 2fr .55fr;gap:clamp(.5rem,1.4vw,1.4rem);align-items:center;font-size:var(--body-size)}.language-row strong{color:var(--teal)}.language-track{height:clamp(.7rem,2vh,1.5rem);background:var(--line)}.language-track div{height:100%;background:var(--teal)}
.architecture{display:grid;gap:clamp(1rem,3vh,3rem)}.pipeline{display:grid;grid-template-columns:.6fr 1fr .15fr 1fr .15fr 1fr .15fr 1.2fr;gap:clamp(.3rem,.7vw,.7rem);align-items:center}.pipeline div,.online>div{background:#FFFDF8;border:1px solid var(--line);padding:clamp(.8rem,1.5vw,1.5rem);font-size:var(--body-size);text-align:center;line-height:1.5}.lane{font-size:var(--small-size);color:var(--teal);font-weight:bold;letter-spacing:.1em}.pipeline small,.online small{font-size:var(--small-size);color:var(--muted)}.pipeline b{color:var(--teal);font-size:var(--body-size)}.pipeline .focus,.online .focus{background:var(--soft);border-color:var(--teal)}.online{display:grid;grid-template-columns:.6fr 1fr 1.4fr 1.4fr 1.2fr;gap:clamp(.5rem,1vw,1rem);align-items:center}.online .retrievers{padding:0;border:0;background:none;display:grid;gap:clamp(.3rem,.7vh,.6rem)}.retrievers div{background:var(--ink);color:var(--paper);padding:clamp(.5rem,1vh,1rem)}.retrievers small{color:#DCE9E6}
.cover{background:var(--ink);color:var(--paper)}.cover::after{content:'A1';position:absolute;right:4vw;top:12vh;font-family:DeckTitle,Georgia,serif;font-size:clamp(5rem,25vw,25rem);line-height:1;opacity:.045;pointer-events:none}.cover .eyebrow{color:#B8D7D0;border-color:#465B66}.cover .slide-content{justify-content:center;gap:clamp(1rem,3vh,3rem)}.cover h1{font-size:clamp(3.2rem,9vw,9rem);line-height:1.1;letter-spacing:-.04em;position:relative;z-index:1}.cover .subtitle{font-size:clamp(1.1rem,3vw,2.5rem);line-height:1.4;max-width:24em;color:#DBE7E3}.cover .institution{font-size:var(--body-size);line-height:1.6;color:#B8C8CB}.team{display:grid;grid-template-columns:repeat(4,1fr);gap:clamp(.7rem,2vw,2rem);border-top:1px solid #465B66;padding-top:clamp(.7rem,2vh,1.5rem);font-size:var(--small-size);line-height:1.6;color:#DBE7E3}.team span{display:block;color:#A4BABD}.cover .source{color:#A4BABD}
.reveal{opacity:0;transform:translateY(8px);transition:opacity 250ms ease,transform 250ms ease}.slide.visible .reveal{opacity:1;transform:none}
.references .card-index{display:none}.references .card p{font-size:clamp(.82rem,1.2vw,1.1rem)}
nav{position:fixed;bottom:clamp(.7rem,1.8vh,1.4rem);left:50%;transform:translateX(-50%);display:flex;gap:clamp(.2rem,.5vw,.4rem);z-index:5;background:#FFFDF8;border:1px solid var(--line);border-radius:30px;padding:clamp(.2rem,.5vw,.5rem)}button{border:0;color:var(--ink);background:transparent;cursor:pointer;padding:clamp(.3rem,.7vw,.7rem);font-size:clamp(.85rem,1.2vw,1rem);font-family:inherit;border-radius:24px}button:hover,button:focus-visible{background:var(--soft);outline:2px solid var(--teal)}#progress{position:fixed;top:0;left:0;height:4px;background:var(--teal);z-index:8;transition:width 150ms}
dialog{border:1px solid var(--line);border-radius:12px;max-width:min(90vw,850px);padding:clamp(1rem,3vw,2rem);background:var(--paper);color:var(--ink)}dialog::backdrop{background:#152C3CCC}dialog h2{font-size:clamp(1.2rem,2.5vw,2rem);margin-bottom:1rem}dialog p{font-size:clamp(.9rem,1.5vw,1.15rem);line-height:1.7;max-height:60vh;overflow:auto}dialog button{float:right}.sr-only{position:absolute;width:1px;height:1px;padding:0;margin:-1px;overflow:hidden;clip:rect(0,0,0,0);white-space:nowrap;border:0}
@media(max-width:600px){:root{--body-size:clamp(.82rem,2.5vw,1rem);--h2-size:clamp(1.45rem,6vw,2rem);--h3-size:clamp(.95rem,3.5vw,1.15rem);--small-size:clamp(.65rem,2vw,.75rem)}.slide-content{gap:clamp(.6rem,1.5vh,1rem);padding-bottom:clamp(3.5rem,9vh,4rem)}.count-3{grid-template-columns:1fr}.count-3 .card{padding:clamp(.6rem,2vw,1rem);display:grid;grid-template-columns:1fr 1.55fr;column-gap:.65rem;align-items:center}.count-3 .card-index{display:none}.count-3 h3{margin:0}.card p{line-height:1.45}.card-index{display:none}.stats{gap:.6rem}.stat strong{font-size:clamp(1.75rem,6vw,2.3rem)}.stat span{font-size:clamp(.75rem,2vw,.9rem)}.steps{grid-template-columns:repeat(2,minmax(0,1fr))}.table-wrap table{font-size:clamp(.72rem,2.2vw,.9rem)}.table-wrap th,.table-wrap td{padding:.7rem .25rem}.experiment{grid-template-columns:1.5rem 1fr 1.2fr;gap:.4rem;font-size:clamp(.75rem,2.1vw,.9rem)}.language-row{grid-template-columns:1fr 1fr .7fr;font-size:clamp(.75rem,2.1vw,.9rem)}.architecture{gap:1rem}.pipeline{grid-template-columns:repeat(4,minmax(0,1fr))}.pipeline .lane{grid-column:1/-1}.pipeline b{display:none}.pipeline div{padding:.65rem .3rem;font-size:clamp(.75rem,2vw,.9rem)}.pipeline small{font-size:clamp(.65rem,1.8vw,.75rem)}.online{grid-template-columns:1fr 1fr}.online .lane{grid-column:1/-1}.online>div{padding:.6rem;font-size:clamp(.75rem,2vw,.9rem)}.team{grid-template-columns:1fr 1fr}.source{font-size:clamp(.6rem,1.8vw,.72rem);bottom:.6rem;max-width:75%}.number{bottom:.6rem}.cover .institution{font-size:clamp(.8rem,2vw,.95rem)}.latency-row{font-size:clamp(.75rem,2.1vw,.9rem);grid-template-columns:7rem 1fr}.takeaway{font-size:clamp(.78rem,2.2vw,.95rem)}}
@media(max-height:600px) and (min-width:601px){:root{--body-size:clamp(.8rem,1.6vw,1rem);--h2-size:clamp(1.3rem,3vw,1.75rem);--h3-size:clamp(.85rem,1.8vw,1.1rem);--small-size:clamp(.65rem,1.1vw,.8rem);--slide-padding:clamp(.7rem,2vw,1.2rem)}.slide-content{gap:.45rem;padding-bottom:2.6rem}.body{gap:.55rem}.card{padding:.6rem .8rem}.card-index{display:none}.card p{line-height:1.4}h3{margin-bottom:.4rem}.stat strong{font-size:clamp(1.8rem,4vw,2.5rem)}.stat span{margin-top:.3rem}.stat{padding-top:.4rem}.steps{margin:0}.steps li{padding:.4rem .55rem}.steps span{margin-bottom:.2rem}.formula{padding:.6rem;font-size:clamp(1.2rem,3vw,1.7rem)}.takeaway{padding:.45rem .65rem}.table-wrap th,.table-wrap td{padding:.35rem .5rem}.experiment{padding:.3rem .6rem}.architecture{gap:.6rem}.pipeline div,.online>div{padding:.5rem}.retrievers div{padding:.35rem}.latency-chart{gap:.5rem}.language-chart{gap:.5rem}.cover h1{font-size:clamp(2.8rem,8vw,4.8rem)}.cover .subtitle{font-size:clamp(1rem,3vw,1.5rem)}.cover .slide-content{gap:.55rem}.team{padding-top:.5rem}.source,.number{bottom:.6rem}nav{bottom:.4rem}}
@media print{@page{size:13.333in 7.5in;margin:0}html,body{height:auto;overflow:visible}html{scroll-snap-type:none}body{print-color-adjust:exact;-webkit-print-color-adjust:exact}.slide{width:13.333in;height:7.5in;break-after:page}.slide-content{padding:40px 50px 60px}.reveal{opacity:1;transform:none}nav,#progress,dialog{display:none!important}.source,.number{bottom:20px}}
'''

JS = r'''
/* Keyboard, wheel and swipe controller; native scrolling still works. */
class PresentationController {
 constructor(){
  this.slides=[...document.querySelectorAll('.slide')];this.current=0;this.lastWheel=0;this.pending=null;
  this.motion=matchMedia('(prefers-reduced-motion: reduce)');
  this.observer=new IntersectionObserver(entries=>{for(const e of entries){if(e.isIntersecting){e.target.classList.add('visible');const index=this.slides.indexOf(e.target);if(this.pending===index&&e.intersectionRatio>.95)this.pending=null;if(this.pending===null&&e.intersectionRatio>.55){this.current=index;this.update();}}}},{threshold:[.15,.55,.95]});
  this.slides.forEach(s=>this.observer.observe(s));
  document.addEventListener('keydown',e=>{
   if(document.querySelector('dialog[open]'))return;
   if(['ArrowRight','ArrowDown','PageDown',' '].includes(e.key)){e.preventDefault();this.go(this.current+1)}
   if(['ArrowLeft','ArrowUp','PageUp'].includes(e.key)){e.preventDefault();this.go(this.current-1)}
   if(e.key==='Home'){e.preventDefault();this.go(0)}if(e.key==='End'){e.preventDefault();this.go(this.slides.length-1)}
   if(e.key.toLowerCase()==='n')this.notes();if(e.key.toLowerCase()==='f')this.fullscreen();
  });
  document.addEventListener('wheel',e=>{if(document.querySelector('dialog[open]')||Math.abs(e.deltaY)<15||Math.abs(e.deltaX)>Math.abs(e.deltaY))return;e.preventDefault();if(Date.now()-this.lastWheel>600){this.lastWheel=Date.now();this.go(this.current+(e.deltaY>0?1:-1))}},{passive:false});
  let start=null;document.addEventListener('touchstart',e=>{if(!document.querySelector('dialog[open]'))start=[e.touches[0].clientX,e.touches[0].clientY]},{passive:true});
  document.addEventListener('touchend',e=>{if(!start)return;let dx=e.changedTouches[0].clientX-start[0],dy=e.changedTouches[0].clientY-start[1];if(Math.abs(dx)>45&&Math.abs(dx)>Math.abs(dy))this.go(this.current+(dx<0?1:-1));start=null},{passive:true});
  document.getElementById('prev').onclick=()=>this.go(this.current-1);document.getElementById('next').onclick=()=>this.go(this.current+1);document.getElementById('notes').onclick=()=>this.notes();document.getElementById('fullscreen').onclick=()=>this.fullscreen();
  document.getElementById('close-notes').onclick=()=>document.getElementById('speaker-notes').close();
  const hash=Number(location.hash.slice(1));this.go(Number.isInteger(hash)&&hash>=1&&hash<=this.slides.length?hash-1:0);this.update();
 }
 go(i){this.current=Math.max(0,Math.min(this.slides.length-1,i));this.pending=this.current;this.slides[this.current].scrollIntoView({behavior:this.motion.matches?'instant':'smooth',block:'start'});if(Math.abs(this.slides[this.current].getBoundingClientRect().top)<1)this.pending=null;this.update()}
 update(){document.getElementById('progress').style.width=((this.current+1)/this.slides.length*100)+'%';document.getElementById('prev').disabled=this.current===0;document.getElementById('next').disabled=this.current===this.slides.length-1;document.getElementById('position').textContent=(this.current+1)+' / '+this.slides.length;history.replaceState(null,'','#'+(this.current+1))}
 notes(){document.getElementById('notes-title').textContent=(this.current+1)+'. '+this.slides[this.current].dataset.title;document.getElementById('notes-text').textContent=NOTES[this.current];document.getElementById('speaker-notes').showModal()}
 async fullscreen(){try{if(document.fullscreenElement)await document.exitFullscreen();else await document.documentElement.requestFullscreen()}catch{document.getElementById('position').textContent='Dùng chế độ toàn màn hình của trình duyệt'}}
}
window.presentation=new PresentationController();
'''

def make_html():
    fonts=''
    for family,filename in [('DeckBody','Verdana.ttf'),('DeckTitle','Georgia Bold.ttf')]:
        p=Path('/System/Library/Fonts/Supplemental')/filename
        if p.exists():
            fonts+=f"@font-face{{font-family:{family};src:url(data:font/ttf;base64,{base64.b64encode(p.read_bytes()).decode()}) format('truetype');font-display:swap;font-weight:{'700' if family=='DeckTitle' else '400'}}}"
    sections=[]
    for i,s in enumerate(SLIDES):
        foot=f'<footer class="source">Nguồn: báo cáo Nhóm 4 · PDF tr. {s["pages"]}</footer><span class="number">{i+1:02} / {len(SLIDES):02}</span>'
        if s['kind']=='cover':
            content=f'''<div class="slide-content"><div class="eyebrow">{s['section']} · NHÓM 4 · THÁNG 9/2026</div><h1 class="reveal">ShopMind</h1><p class="subtitle reveal">{s['subtitle']}</p><div class="institution">Môn: Tìm kiếm thông minh · UIT, ĐHQG TP.HCM<br>GVHD: TS. Phạm Đỗ Kim Chi</div><div class="team"><div>Trần Văn Quang<span>250104022</span></div><div>Nguyễn Minh Thái<span>250201090</span></div><div>Trần Khoa<span>250104010</span></div><div>Đinh Ngọc Phúc<span>26848170</span></div></div></div>'''
        else:
            takeaway='<div class="takeaway reveal">'+esc(s['takeaway'])+'</div>' if 'takeaway' in s else ''
            content=f'<div class="slide-content"><div class="eyebrow">{esc(s["section"])}</div><h2 class="reveal">{esc(s["title"])}</h2><p class="lead reveal">{esc(s["lead"])}</p><div class="body reveal">{body_html(s)}</div>{takeaway}</div>'
        sections.append(f'<section class="slide {s["kind"]}" id="slide-{i+1}" aria-label="Slide {i+1}: {escape(s["title"],quote=True)}" data-title="{escape(s["title"],quote=True)}">{content}{foot}</section>')
    notes=[s['notes']+f'\n\nNguồn: {SOURCE}, PDF trang {s["pages"]}.' for s in SLIDES]
    html='<!doctype html><html lang="vi"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>ShopMind A1 · Nhóm 4</title><style>'+fonts+CSS+'</style></head><body><main>'+''.join(sections)+'</main><div id="progress" aria-hidden="true"></div><nav aria-label="Điều khiển trình chiếu"><button id="prev" aria-label="Slide trước">←</button><span class="sr-only" id="position" aria-live="polite"></span><button id="next" aria-label="Slide tiếp">→</button><button id="notes" aria-label="Ghi chú thuyết trình (N)">Ghi chú</button><button id="fullscreen" aria-label="Toàn màn hình (F)">⛶</button></nav><dialog id="speaker-notes" aria-labelledby="notes-title"><button id="close-notes" aria-label="Đóng ghi chú">Đóng ×</button><h2 id="notes-title"></h2><p id="notes-text"></p></dialog><script>const NOTES='+json.dumps(notes,ensure_ascii=False)+';'+JS+'</script></body></html>'
    (OUT/'ShopMind_A1.html').write_text(html,encoding='utf-8')

COLORS=dict(paper='F5F3ED',ink='152C3C',muted='536774',teal='086F69',coral='B44B34',line='D7DFDA',soft='E6EDE8',white='FFFDF8',p95='A9C8C5')
def color(c): return RGBColor.from_string(COLORS.get(c,c))

def rect(sl,x,y,w,h,fill='white',border=None):
    shape=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    shape.fill.solid();shape.fill.fore_color.rgb=color(fill)
    if border:shape.line.color.rgb=color(border);shape.line.width=Pt(.7)
    else:shape.line.fill.background()
    return shape

def text(sl,x,y,w,h,string,size=17,c='ink',bold=False,font='Verdana',align=None):
    shape=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=shape.text_frame;tf.word_wrap=True
    tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
    for i,line in enumerate(str(string).split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line;p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=color(c)
        p.space_after=Pt(5);p.space_before=Pt(0);p.line_spacing=1.2
        if align is not None:p.alignment=align
    return shape

def card(sl,x,y,w,h,heading,body,num=None,size=16):
    rect(sl,x,y,w,h,'white','line');rect(sl,x,y,w,.035,'teal')
    if h<2:
        text(sl,x+.22,y+.15,w-.44,.35,heading,17,'ink',True)
        shape=text(sl,x+.22,y+.59,w-.44,h-.7,body,14,'muted')
        for p in shape.text_frame.paragraphs:
            p.line_spacing=1.12;p.space_after=Pt(0)
        return
    if num is not None:text(sl,x+.22,y+.18,w-.44,.24,f'{num:02}',10,'teal',True)
    title_y=y+(.55 if num is not None else .2)
    text(sl,x+.22,title_y,w-.44,.5,heading,19,'ink',True)
    text(sl,x+.22,title_y+.66,w-.44,h-(title_y-y)-.77,body,size,'muted')

def card_grid(sl,items,y=2.6,h=2.95):
    n=len(items);cols=3 if n==3 else 2;rows=2 if n==4 else 1
    gap=.2;width=(12-cols*gap+gap)/cols;height=(h-(rows-1)*gap)/rows
    for i,(heading,body) in enumerate(items):
        card(sl,.66+(i%cols)*(width+gap),y+(i//cols)*(height+gap),width,height,heading,body,i+1,size=14.5 if n==4 else 16)

def stats(sl,items,y=2.7,h=1.45):
    width=12/len(items)
    for i,(value,label) in enumerate(items):
        x=.66+i*width;rect(sl,x,y,width-.2,.015,'line')
        text(sl,x,y+.16,width-.2,.75,value,44 if len(items)<4 else 40,'teal',True,'Georgia')
        text(sl,x,y+.99,width-.2,.55,label,15,'muted')

def steps(sl,items,y=4.8):
    w=(12-.45)/4
    for i,s in enumerate(items):
        x=.66+i*(w+.15);rect(sl,x,y,w,.85,'soft')
        text(sl,x+.16,y+.12,w-.32,.18,f'{i+1:02}',10,'teal',True)
        text(sl,x+.16,y+.38,w-.32,.4,s,13,'ink')

def make_pptx():
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    prs.core_properties.title='ShopMind – Báo cáo A1 · Nhóm 4'
    prs.core_properties.subject='Tìm kiếm sản phẩm đa phương thức'
    prs.core_properties.author='Nhóm 4';prs.core_properties.language='vi-VN'
    prs.core_properties.comments='Nội dung tổng hợp từ Nhom4_BaoCao-A1.pdf; kết quả không được chạy lại khi tạo slide.'
    for i,s in enumerate(SLIDES):
        sl=prs.slides.add_slide(prs.slide_layouts[6]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=color('ink' if i==0 else 'paper')
        sl.notes_slide.notes_text_frame.text=s['notes']+f'\n\nNguồn: {SOURCE}, trang PDF {s["pages"]} (tính cả bìa/mục lục).'
        if i==0:
            text(sl,.68,.58,11.7,.35,'BÁO CÁO BÀI TẬP 1  /  NHÓM 4  /  THÁNG 9/2026',12,'A9C8C5',True)
            rect(sl,.68,1.02,11.97,.014,'536774')
            text(sl,.68,1.7,11.4,1.5,'ShopMind',82,'paper',True,'Georgia')
            text(sl,.72,3.3,11.6,.7,s['subtitle'],30,'D7DFDA')
            text(sl,.72,4.45,11.6,.7,'Môn: Tìm kiếm thông minh · UIT, ĐHQG TP.HCM\nGVHD: TS. Phạm Đỗ Kim Chi',16,'A9C8C5')
            rect(sl,.72,5.62,11.9,.014,'536774')
            for j,(name,ident) in enumerate([('Trần Văn Quang','250104022'),('Nguyễn Minh Thái','250201090'),('Trần Khoa','250104010'),('Đinh Ngọc Phúc','26848170')]):
                text(sl,.72+j*3,5.9,2.9,.7,name+'\n'+ident,13,'D7DFDA')
        else:
            text(sl,.66,.42,12,.27,s['section'],10.5,'teal',True)
            rect(sl,.66,.89,12,.014,'line')
            text(sl,.66,1.1,12,1.06,s['title'],32 if len(s['title'])<47 else 29,'ink',True,'Georgia')
            text(sl,.66,2.05,12,.47,s['lead'],14,'muted')
            k=s['kind']
            if k in ('problem','benchmark','modes','dense','limits','conclusion','references'):
                card_grid(sl,s['cards'],y=2.78,h=3.04)
            elif k=='architecture':
                for y,label,items in [(2.9,'OFFLINE',[('ABO','metadata + ảnh'),('Chuẩn hóa','catalog chung'),('SigLIP2','text + image vectors'),('Elasticsearch','metadata + vectors')]),(4.6,'ONLINE',[('Text / ảnh','truy vấn người dùng'),('BM25 / kNN','text_vector / image_vector'),('Độc lập / RRF','Hybrid: BM25 + dense text'),('API → top-k','danh sách sản phẩm')])]:
                    text(sl,.66,y-.3,12,.2,label,10,'teal',True)
                    w=2.72
                    for j,(a,b) in enumerate(items):
                        x=.66+j*3.09;rect(sl,x,y,w,1.05,'soft' if j==3 else 'white','teal' if j==3 else 'line')
                        text(sl,x+.12,y+.2,w-.24,.3,a,18,'teal' if j==3 else 'ink',True,align=PP_ALIGN.CENTER)
                        text(sl,x+.1,y+.65,w-.2,.27,b,11.2,'muted',align=PP_ALIGN.CENTER)
                        if j<3:text(sl,x+w+.08,y+.35,.3,.3,'→',17,'teal')
            elif k in ('dataset','image'):
                stats(sl,s['stats'],y=2.85);steps(sl,s['steps'],y=4.85)
            elif k=='fusion':
                rect(sl,.66,2.7,12,.95,'ink');text(sl,.85,2.87,11.6,.6,s['formula'],30,'paper',False,'Georgia',PP_ALIGN.CENTER)
                for j,(a,b) in enumerate(s['cards']):card(sl,.66+j*6.1,3.88,5.9,1.97,a,b,size=16)
            elif k=='experiments':
                for j,(a,b,c) in enumerate(s['rows']):
                    y=2.75+j*.51;rect(sl,.66,y,12,.48,'soft' if j%2==0 else 'paper')
                    text(sl,.83,y+.08,.65,.32,a,14,'teal',True);text(sl,1.65,y+.08,3.3,.32,b,14,'ink',True);text(sl,5.3,y+.08,7.1,.32,c,14,'muted')
            elif k=='results':
                x=[.66,6.8,8.8,10.8];w=[6.1,2,2,1.86]
                rect(sl,.66,2.78,12,.53,'ink')
                for j,a in enumerate(['Cấu hình','Recall@10','nDCG@10','MRR@10']):text(sl,x[j]+.15,2.94,w[j]-.3,.25,a,14,'paper',True)
                for j,r in enumerate(TEXT_RESULTS):
                    y=3.31+j*.48;rect(sl,.66,y,12,.48,'soft' if j==0 else ('white' if j%2 else 'paper'))
                    for col,a in enumerate([r[0]]+[comma(v) for v in r[1:4]]):text(sl,x[col]+.15,y+.1,w[col]-.3,.3,a,16,'teal' if j==0 else 'ink',j==0)
            elif k=='latency':
                text(sl,3.6,2.65,9,.22,'p50: teal     p95: xanh nhạt                 Trục chung 0–120 ms',11,'muted')
                for j,r in enumerate(TEXT_RESULTS):
                    y=3.08+j*.55;text(sl,.68,y+.08,2.8,.38,r[0],16)
                    for offset,val,fill in [(0,r[4],'teal'),(.23,r[5],'p95')]:
                        barw=7.8*val/120;rect(sl,3.6,y+offset,barw,.17,fill);text(sl,3.72+barw,y+offset-.025,.85,.22,comma(val,2),10.5)
            elif k=='language':
                for j,(a,v) in enumerate(s['bars']):
                    y=2.85+j*.7;text(sl,.68,y,3.7,.4,a,17)
                    rect(sl,4.35,y+.03,6.5,.28,'line');rect(sl,4.35,y+.03,6.5*v,.28,'teal');text(sl,11.05,y,1.5,.4,comma(v*100,2)+'%',20,'teal',True)
                for j,(a,b) in enumerate(s['cards']):card(sl,.66+j*6.1,4.53,5.9,1.75,a,b,size=15)
            elif k in ('ablation','id'):
                stats(sl,s['stats'],y=2.75)
                for j,(a,b) in enumerate(s['cards']):card(sl,.66+j*6.1,4.6,5.9,1.4,a,b,size=15)
            if 'takeaway' in s:
                rect(sl,.66,6.25,12,.57,'soft');rect(sl,.66,6.25,.035,.57,'teal')
                text(sl,.83,6.38,11.65,.37,s['takeaway'],12.5,'teal')
        text(sl,.66,7.08,10.8,.24,f'Nguồn: báo cáo Nhóm 4 · PDF tr. {s["pages"]}',9.5,'A9C8C5' if i==0 else 'muted')
        text(sl,11.7,7.08,.97,.24,f'{i+1:02} / {len(SLIDES):02}',9.5,'A9C8C5' if i==0 else 'ink',align=PP_ALIGN.RIGHT)
    prs.save(OUT/'ShopMind_A1.pptx')

def make_notes():
    lines=['# ShopMind A1 – Ghi chú thuyết trình','',f'Nguồn chính: `{SOURCE}`. Trang nguồn tính theo thứ tự trang PDF, gồm bìa và mục lục.','', '18 slide · tiếng Việt · thời lượng gợi ý 15–18 phút. Số liệu giữ theo báo cáo; không chạy lại benchmark khi tạo slide.','']
    for i,s in enumerate(SLIDES):
        lines += [f'## {i+1:02}. {s["title"]}', '', f'Nguồn: PDF trang {s["pages"]}.', '',s['notes'],'']
    (OUT/'SPEAKER_NOTES.md').write_text('\n'.join(lines),encoding='utf-8')
    (OUT/'content.json').write_text(json.dumps(dict(source=SOURCE,slides=SLIDES,text_results=TEXT_RESULTS),ensure_ascii=False,indent=2)+'\n',encoding='utf-8')

if __name__=='__main__':
    make_html();make_pptx();make_notes()
    print(f'Created {len(SLIDES)} slides in {OUT}')
